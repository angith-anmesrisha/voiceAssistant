"""PowerPoint decks: request parsing, the free-model writer, the JSON clean-up, drawing the
.pptx, and the tool with an uploaded file."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pytest

from jarvis.agent.assistant import Assistant
from jarvis.app.config import ToolSettings
from jarvis.app.events import ChatEntry, EntryKind, EventBus
from jarvis.app.state import StateMachine
from jarvis.providers.base import ChatResult, ErrorKind, ProviderError
from jarvis.storage.db import HistoryStore
from jarvis.tools.apps import AppIndex
from jarvis.tools.builtin import SlideTools, build_registry
from jarvis.tools.deck_render import render
from jarvis.tools.files import FileArea
from jarvis.tools.slides import (
    DeckError,
    DeckRequest,
    DeckWriter,
    deck_command,
    excerpt,
    is_free,
    parse_deck,
    unique_path,
)
from tests.fakes import FakeDesktop
from tests.test_assistant import ScriptedRouter


def deck_json(count: int = 6, theme: str = "ocean") -> str:
    slides: list[dict[str, Any]] = [
        {"layout": "title", "title": "Black holes", "subtitle": "What we know", "notes": "Hi."}
    ]
    middle: list[dict[str, Any]] = [
        {"layout": "section", "title": "The basics"},
        {"layout": "bullets", "title": "Gravity wins", "bullets": ["Mass collapses", "Light"]},
        {"layout": "stats", "title": "By the numbers",
         "stats": [{"value": "1 in 10", "label": "galaxies"}]},
        {"layout": "timeline", "title": "History",
         "steps": [{"label": "1916", "text": "Schwarzschild"}, {"label": "2019", "text": "Image"}]},
        {"layout": "comparison", "title": "Kinds", "columns": ["", "Stellar", "Supermassive"],
         "rows": [["Mass", "10 suns", "Millions"]]},
        {"layout": "cards", "title": "Four facts", "bullets": ["a", "b", "c", "d"]},
        {"layout": "two_column", "title": "Myths and facts",
         "left": {"heading": "Myth", "bullets": ["They suck"]},
         "right": {"heading": "Fact", "bullets": ["They pull"]}},
        {"layout": "quote", "quote": "Black holes ain't so black.", "attribution": "Hawking"},
    ]  # fmt: skip
    while len(slides) < count - 1:
        slides.append(dict(middle[(len(slides) - 1) % len(middle)]))
    slides.append({"layout": "closing", "title": "Takeaways", "bullets": ["Strange", "Real"]})
    return json.dumps({"title": "Black Holes", "subtitle": "", "theme": theme, "slides": slides})


# --- what the user says ---


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("make a 20 slide ppt for the pdf uploaded", DeckRequest("", 20, "upload")),
        ("Make a twenty slide presentation on the pdf I'll upload", DeckRequest("", 20, "upload")),
        ("create a presentation on black holes", DeckRequest("black holes", 10, "")),
        ("make a ppt about climate change in 12 slides", DeckRequest("climate change", 12, "")),
        (
            "build a 15-slide deck about India's space programme",
            DeckRequest("India's space programme", 15, ""),
        ),
        ("make slides from the document I'm attaching", DeckRequest("", 10, "upload")),
    ],
)
def test_deck_requests_are_understood(text: str, expected: DeckRequest) -> None:
    assert deck_command(text) == expected


def test_other_requests_are_left_to_the_model() -> None:
    assert deck_command("open powerpoint") is None
    assert deck_command("make a ppt from report.pdf in Downloads") is None
    assert deck_command("what's on my screen") is None


# --- the model's answer ---


def test_deck_json_is_cleaned_up() -> None:
    reply = "<think>plan</think>```json\n" + deck_json(10) + "\n```"
    deck = parse_deck(reply, 10)
    assert deck.theme == "ocean" and len(deck.slides) == 10
    assert deck.slides[0].layout == "title" and deck.slides[-1].layout == "closing"
    table = next(s for s in deck.slides if s.layout == "comparison")
    assert table.columns == ["", "Stellar", "Supermassive"]  # the blank corner stays


def test_layouts_without_their_content_become_bullets() -> None:
    reply = json.dumps(
        {
            "title": "T",
            "theme": "neon",
            "slides": [
                {"layout": "stats", "title": "No numbers", "bullets": ["a"]},
                {"layout": "timeline", "title": "One step", "steps": [{"label": "x", "text": "y"}]},
                {"layout": "hologram", "title": "Odd", "bullets": ["b"]},
            ],
        }
    )
    deck = parse_deck(reply, 3)
    assert [s.layout for s in deck.slides] == ["bullets", "bullets", "bullets"]
    assert deck.theme == "midnight"
    assert deck.slides[1].bullets == ["x: y"]


@pytest.mark.parametrize("reply", ["no json here", '{"slides": []}', "{broken", deck_json(3)])
def test_unusable_answers_are_rejected(reply: str) -> None:
    with pytest.raises(DeckError):
        parse_deck(reply, 20)


def test_long_documents_are_sampled_end_to_end() -> None:
    text = "".join(f"Section {i}. " + "words " * 400 for i in range(100))
    sample, cut = excerpt(text, 12_000)
    assert cut and len(sample) <= 12_500
    assert "Section 0." in sample and "Section 9" in sample[-2000:]
    assert excerpt("short") == ("short", False)


# --- the writer: free models only ---


class FakeProvider:
    def __init__(self, name: str, replies: list[Any], key: bool = True) -> None:
        self.display_name = name
        self.replies = replies
        self.key = key
        self.calls: list[tuple[str, int]] = []

    def missing_config(self) -> str | None:
        return None if self.key else "no key"

    async def complete(self, model: str, messages: list[Any], **kwargs: Any) -> ChatResult:
        self.calls.append((model, sum(len(m["content"]) for m in messages)))
        reply = self.replies.pop(0)
        if isinstance(reply, Exception):
            raise reply
        return ChatResult(self.display_name, model, reply, [], {})


def test_only_free_models_are_ever_used() -> None:
    assert is_free("qwen/qwen3.8-27b:free") and is_free("openrouter/free")
    assert not is_free("moonshotai/kimi-k3")


async def test_the_next_free_model_is_tried_when_one_fails() -> None:
    busy = ProviderError("OpenRouter", ErrorKind.QUOTA, "HTTP 429: rate limited")
    openrouter = FakeProvider("OpenRouter", [busy, "not json", deck_json(5)])
    writer = DeckWriter(
        openrouter,
        None,
        ["a/one:free", "moonshotai/kimi-k3", "b/two:free", "c/three:free"],
        groq_model="",
    )
    written = await writer.write("black holes", 5)
    assert [m for m, _ in openrouter.calls] == ["a/one:free", "b/two:free", "c/three:free"]
    assert (written.provider, written.model) == ("OpenRouter", "c/three:free")


async def test_groq_is_the_free_backup_for_short_requests_only() -> None:
    down = ProviderError("OpenRouter", ErrorKind.UNAVAILABLE, "HTTP 503")
    groq = FakeProvider("Groq", [deck_json(4)])
    writer = DeckWriter(FakeProvider("OpenRouter", [down]), groq, ["a/one:free"])
    written = await writer.write("tea", 4)
    assert written.model == "openai/gpt-oss-120b" and groq.calls
    long_doc = "x" * 20_000
    with pytest.raises(DeckError, match="too long for its free tier"):
        await DeckWriter(FakeProvider("OpenRouter", [down]), groq, ["a/one:free"]).write(
            "", 4, long_doc, "big.pdf"
        )


async def test_no_openrouter_key_says_how_to_get_one() -> None:
    writer = DeckWriter(FakeProvider("OpenRouter", [], key=False), None, ["a:free"])
    with pytest.raises(DeckError, match="OPENROUTER_API_KEY"):
        await writer.write("tea", 5)


# --- the file ---


def test_every_layout_draws_and_carries_notes(tmp_path: Path) -> None:
    from pptx import Presentation

    for theme in ("midnight", "daylight"):
        path = tmp_path / f"{theme}.pptx"
        assert render(parse_deck(deck_json(12, theme), 12), path) == 12
        prs = Presentation(str(path))
        assert len(prs.slides) == 12
        assert prs.slides[0].notes_slide.notes_text_frame.text == "Hi."
        text = " ".join(
            shape.text_frame.text
            for slide in prs.slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
        assert "Black holes" in text and "Hawking" in text
    with pytest.raises(FileExistsError):
        render(parse_deck(deck_json(4), 4), tmp_path / "midnight.pptx")  # never overwrites


def test_file_names_are_safe_and_new(tmp_path: Path) -> None:
    first = unique_path(tmp_path, 'AI: "the future"?')
    assert first.name == "AI the future.pptx"
    first.write_bytes(b"x")
    second = unique_path(tmp_path, 'AI: "the future"?')
    assert second != first and second.suffix == ".pptx"


# --- the tool and the assistant ---


def slide_registry(tmp_path: Path, reply: str, picked: Path | None) -> tuple[Any, Any, list[Path]]:
    openrouter = FakeProvider("OpenRouter", [reply])
    opened: list[Path] = []
    desktop = FakeDesktop()
    desktop.open_path = opened.append  # type: ignore[method-assign,assignment]

    async def pick() -> Path | None:
        return picked

    registry = build_registry(
        ToolSettings(),
        desktop,
        AppIndex([], apps_folder=False),
        files=FileArea([tmp_path]),
        slides=SlideTools(
            DeckWriter(openrouter, None, ["best/model:free"], groq_model=""),
            tmp_path / "Presentations",
            pick_file=pick,
        ),
    )
    return registry, openrouter, opened


async def test_a_deck_from_an_uploaded_file(tmp_path: Path) -> None:
    notes = tmp_path / "notes.md"
    notes.write_text("# Photosynthesis\nPlants turn light into sugar.", encoding="utf-8")
    registry, openrouter, opened = slide_registry(tmp_path, deck_json(5), notes)
    spec = registry.get("make_presentation")
    assert spec is not None
    result = await registry.run(spec, {"slides": 5, "source_file": "upload"})
    assert result.ok, result.summary
    assert "best/model:free" not in result.summary and "model (OpenRouter, free)" in result.summary
    assert opened and opened[0].suffix == ".pptx" and opened[0].exists()
    assert "Plants turn light into sugar" in str(openrouter.calls) or openrouter.calls[0][1] > 100
    nothing, _, _ = slide_registry(tmp_path, deck_json(5), None)
    spec = nothing.get("make_presentation")
    assert spec is not None
    declined = await nothing.run(spec, {"source_file": "upload"})
    assert not declined.ok and "No file was picked" in declined.summary


async def test_assistant_makes_a_deck_without_asking_the_chat_model(tmp_path: Path) -> None:
    registry, _, _ = slide_registry(tmp_path, deck_json(6), None)
    bus = EventBus()
    entries: list[ChatEntry] = []

    async def on_entry(entry: ChatEntry) -> None:
        entries.append(entry)

    bus.subscribe(ChatEntry, on_entry)

    async def confirm(_: Any) -> bool:
        return True

    async def capture() -> bytes | None:
        return None

    router = ScriptedRouter()  # no chat replies queued: a model call would fail
    assistant = Assistant(
        router,  # type: ignore[arg-type]
        HistoryStore(tmp_path / "h.sqlite3"),
        registry,
        bus,
        StateMachine(bus),
        confirm,
        capture,
    )
    reply = await assistant.handle("make a 6 slide presentation on black holes")
    assert reply.ok and "6 slides" in reply.text and router.requests == []
    last = [e for e in entries if e.kind is EntryKind.ASSISTANT][-1]
    assert (last.provider, last.model, last.mode) == ("OpenRouter", "best/model:free", "slides")
    assert list((tmp_path / "Presentations").glob("*.pptx"))
