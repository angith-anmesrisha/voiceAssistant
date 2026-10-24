"""Drawing a slide deck as a real PowerPoint file (python-pptx, MIT).

The writer model only decides the words, numbers, icons and each slide's layout;
everything visual is done here, the same way every time: 16:9 slides on a theme colour,
two accent colours, large type, lots of room, and layouts drawn from shapes (icon cards,
stat tiles, a timeline, a comparison table, native charts PowerPoint can edit) instead of
walls of bullets. Bullets written as "Lead: detail" get a bold lead. Icons are emoji in
Segoe UI Emoji, and the text is Segoe UI, both on every Windows PC. Every slide gets the
speaker notes the model wrote.
"""

from __future__ import annotations

import math
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

LAYOUTS = (
    "title",
    "agenda",
    "section",
    "bullets",
    "cards",
    "two_column",
    "stats",
    "quote",
    "timeline",
    "comparison",
    "chart",
    "statement",
    "closing",
)

CHART_TYPES = ("column", "bar", "line", "pie", "doughnut")


@dataclass(frozen=True, slots=True)
class Theme:
    background: str
    surface: str
    text: str
    muted: str
    accent: str
    accent2: str
    dark: bool = True


THEMES = {
    "midnight": Theme("0F172A", "1E293B", "F8FAFC", "94A3B8", "38BDF8", "F59E0B"),
    "ocean": Theme("062C30", "0B4A50", "F0FDFA", "8FC9C4", "2DD4BF", "FDE68A"),
    "ember": Theme("1C1917", "2C2724", "FAFAF9", "A8A29E", "F97316", "FACC15"),
    "royal": Theme("1E1B4B", "2E2A6B", "EEF2FF", "A5B4FC", "A78BFA", "F472B6"),
    "daylight": Theme("FFFFFF", "F1F5F9", "0F172A", "475569", "2563EB", "F97316", dark=False),
    "forest": Theme("F7F5EF", "E8E4D8", "1F2A1F", "5B6B5B", "2F6B3B", "C2410C", dark=False),
}

TITLE_FONT = "Segoe UI Semibold"
BODY_FONT = "Segoe UI"
ICON_FONT = "Segoe UI Emoji"


@dataclass
class Chart:
    kind: str  # one of CHART_TYPES
    categories: list[str]
    series: list[tuple[str, list[float]]]  # (name, one value per category)
    unit: str = ""


@dataclass
class Slide:
    layout: str
    title: str = ""
    subtitle: str = ""
    bullets: list[str] = field(default_factory=list)
    left: tuple[str, list[str]] = ("", [])
    right: tuple[str, list[str]] = ("", [])
    stats: list[tuple[str, str]] = field(default_factory=list)  # (value, label)
    quote: str = ""
    attribution: str = ""
    steps: list[tuple[str, str]] = field(default_factory=list)  # (label, text)
    columns: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    takeaway: str = ""
    notes: str = ""
    icon: str = ""
    icons: list[str] = field(default_factory=list)  # per card or stat
    chart: Chart | None = None


@dataclass
class Deck:
    title: str
    subtitle: str
    theme: str
    slides: list[Slide]


# --- geometry helpers (inches) ---

W, H = 13.333, 7.5
MARGIN = 0.75


def text_height(text: str, size: float, width: float) -> float:
    """Roughly how tall (inches) ``text`` wraps to at ``size`` pt in ``width`` inches."""
    per_line = max(1, int(width * 72 / (size * 0.52)))
    lines = sum(max(1, math.ceil(len(part) / per_line)) for part in text.split("\n"))
    return lines * size / 72 * 1.2 + 0.15


def mix(first: str, second: str, amount: float) -> str:
    """``first`` moved ``amount`` (0-1) of the way to ``second`` (hex colours)."""
    a, b = bytes.fromhex(first), bytes.fromhex(second)
    return bytes(round(x + (y - x) * amount) for x, y in zip(a, b, strict=True)).hex().upper()


def split_lead(item: str) -> tuple[str, str]:
    """("Lead", "detail") for "Lead: detail" with a short lead, else ("", item)."""
    head, sep, rest = item.partition(": ")
    if sep and rest and len(head) <= 42 and len(head.split()) <= 6:
        return head, rest
    return "", item


def fitted(text: str, base: float, smallest: float, room: int) -> float:
    """A font size (pt) that lets ``text`` fit about ``room`` characters at ``base``."""
    length = max(len(text), 1)
    if length <= room:
        return base
    return max(smallest, round(base * math.sqrt(room / length), 1))


class _Painter:
    def __init__(self, theme: Theme, deck_title: str) -> None:
        from pptx import Presentation
        from pptx.util import Inches

        self.theme = theme
        self.deck_title = deck_title
        self.prs = Presentation()
        self.prs.slide_width = Inches(W)
        self.prs.slide_height = Inches(H)
        self._blank = self.prs.slide_layouts[6]
        self.number = 0

    # --- primitives ---

    def rgb(self, hex_colour: str) -> Any:
        from pptx.dml.color import RGBColor

        return RGBColor(*bytes.fromhex(hex_colour))  # type: ignore[no-untyped-call]

    def new_slide(self, notes: str, *, footer: bool = True) -> Any:
        slide = self.prs.slides.add_slide(self._blank)
        self.number += 1
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = self.rgb(self.theme.background)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        if footer:
            self.text(slide, self.deck_title, MARGIN, H - 0.5, 8, 0.3, 10, self.theme.muted)
            self.text(
                slide,
                str(self.number),
                W - MARGIN - 1,
                H - 0.5,
                1,
                0.3,
                10,
                self.theme.muted,
                align="right",
            )
        return slide

    def box(
        self,
        slide: Any,
        x: float,
        y: float,
        w: float,
        h: float,
        colour: str,
        *,
        shape: str = "rect",
        line: str = "",
    ) -> Any:
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches, Pt

        kind = {
            "rect": MSO_SHAPE.RECTANGLE,
            "round": MSO_SHAPE.ROUNDED_RECTANGLE,
            "oval": MSO_SHAPE.OVAL,
        }[shape]
        item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        item.fill.solid()
        item.fill.fore_color.rgb = self.rgb(colour)
        if line:
            item.line.color.rgb = self.rgb(line)
            item.line.width = Pt(1)
        else:
            item.line.fill.background()
        item.shadow.inherit = False
        if shape == "round":
            item.adjustments[0] = 0.08
        return item

    def text(
        self,
        slide: Any,
        text: str,
        x: float,
        y: float,
        w: float,
        h: float,
        size: float,
        colour: str,
        *,
        bold: bool = False,
        font: str = BODY_FONT,
        align: str = "left",
        anchor: str = "top",
        italic: bool = False,
    ) -> Any:
        from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
        from pptx.util import Inches, Pt

        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.margin_left = frame.margin_right = Inches(0.05)
        frame.margin_top = frame.margin_bottom = Inches(0.03)
        frame.vertical_anchor = {
            "top": MSO_ANCHOR.TOP,
            "middle": MSO_ANCHOR.MIDDLE,
            "bottom": MSO_ANCHOR.BOTTOM,
        }[anchor]
        for index, line in enumerate(text.split("\n")):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = {
                "left": PP_ALIGN.LEFT,
                "center": PP_ALIGN.CENTER,
                "right": PP_ALIGN.RIGHT,
            }[align]
            run = paragraph.add_run()
            run.text = line
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.name = font
            run.font.color.rgb = self.rgb(colour)
        return box

    def rich(
        self,
        slide: Any,
        parts: list[tuple[str, str, bool]],
        x: float,
        y: float,
        w: float,
        h: float,
        size: float,
        *,
        align: str = "left",
        anchor: str = "top",
    ) -> Any:
        """One paragraph of (text, colour, bold) runs."""
        from pptx.util import Pt

        box = self.text(slide, "", x, y, w, h, size, self.theme.text, align=align, anchor=anchor)
        paragraph = box.text_frame.paragraphs[0]
        for index, (text, colour, bold) in enumerate(parts):
            run = paragraph.runs[0] if index == 0 else paragraph.add_run()
            run.text = text
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.name = TITLE_FONT if bold else BODY_FONT
            run.font.color.rgb = self.rgb(colour)
        return box

    def lead_parts(self, item: str, lead_colour: str = "") -> list[tuple[str, str, bool]]:
        head, rest = split_lead(item)
        if not head:
            return [(item, self.theme.text, False)]
        return [
            (head + ": ", lead_colour or self.theme.accent, True),
            (rest, self.theme.text, False),
        ]

    def icon(
        self, slide: Any, icon: str, x: float, y: float, size: float, *, ring: str = ""
    ) -> None:
        """An emoji in a round badge, ``size`` inches across."""
        if not icon:
            return
        self.box(slide, x, y, size, size, ring or self.theme.surface, shape="oval")
        self.text(
            slide, icon, x, y, size, size, size * 30, self.theme.text,
            font=ICON_FONT, align="center", anchor="middle",
        )  # fmt: skip

    def bullet_list(
        self,
        slide: Any,
        items: list[str],
        x: float,
        y: float,
        w: float,
        h: float,
        size: float,
        *,
        marker: str | None = None,
        smallest: float = 12,
    ) -> None:
        """Items with a small accent square, each as tall as its text, shrunk to fit."""
        if not items:
            return
        marker = marker or self.theme.accent
        while True:
            heights = [text_height(item, size, w - 0.3) for item in items]
            if sum(heights) + 0.12 * len(items) <= h or size <= smallest:
                break
            size = max(smallest, size - 1)
        gap = min(max((h - sum(heights)) / len(items), 0.12), size / 72 * 1.1)
        top = y
        for item, height in zip(items, heights, strict=True):
            self.box(slide, x, top + size / 72 * 0.42, 0.12, 0.12, marker)
            self.rich(slide, self.lead_parts(item), x + 0.3, top, w - 0.3, height, size)
            top += height + gap

    def heading(self, slide: Any, title: str, subtitle: str = "") -> float:
        """The slide title with an accent rule; returns where the body can start."""
        size = fitted(title, 34, 24, 48)
        self.text(
            slide,
            title,
            MARGIN,
            0.45,
            W - 2 * MARGIN,
            0.9,
            size,
            self.theme.text,
            font=TITLE_FONT,
            anchor="bottom",
        )
        self.box(slide, MARGIN + 0.05, 1.42, 0.9, 0.06, self.theme.accent)
        top = 1.7
        if subtitle:
            self.text(slide, subtitle, MARGIN, 1.6, W - 2 * MARGIN, 0.5, 16, self.theme.muted)
            top = 2.2
        return top

    def takeaway(self, slide: Any, text: str) -> None:
        if not text:
            return
        self.box(slide, MARGIN, H - 1.35, W - 2 * MARGIN, 0.65, self.theme.surface, shape="round")
        self.box(slide, MARGIN, H - 1.35, 0.1, 0.65, self.theme.accent2)
        self.text(
            slide,
            text,
            MARGIN + 0.3,
            H - 1.35,
            W - 2 * MARGIN - 0.5,
            0.65,
            fitted(text, 16, 12, 110),
            self.theme.text,
            anchor="middle",
            bold=True,
        )

    # --- layouts ---

    def title_slide(self, s: Slide) -> None:
        slide = self.new_slide(s.notes, footer=False)
        t = self.theme
        self.box(slide, W - 4.6, -1.4, 6.2, 6.2, t.surface, shape="oval")
        self.box(slide, W - 2.6, 4.3, 3.6, 3.6, t.accent, shape="oval")
        title = s.title or self.deck_title
        size = fitted(title, 48, 30, 40)
        title_h = text_height(title, size, 8.2)
        sub_h = text_height(s.subtitle, 20, 8.0) if s.subtitle else 0.0
        top = 3.5 - (title_h + sub_h) / 2
        self.box(slide, MARGIN, top + 0.1, 0.12, title_h + sub_h - 0.1, t.accent)
        self.text(slide, title, MARGIN + 0.4, top, 8.2, title_h, size, t.text, font=TITLE_FONT)
        if s.subtitle:
            self.text(
                slide,
                s.subtitle,
                MARGIN + 0.4,
                top + title_h,
                8.0,
                sub_h,
                fitted(s.subtitle, 20, 14, 90),
                t.muted,
            )

    def section_slide(self, s: Slide, index: int) -> None:
        slide = self.new_slide(s.notes)
        t = self.theme
        self.box(slide, 0, 0, 4.4, H, t.accent)
        self.text(
            slide, f"{index:02d}", 0.6, 2.6, 3.2, 1.8, 96, t.background, font=TITLE_FONT, bold=True
        )
        if s.icon:
            self.icon(slide, s.icon, 0.75, 4.55, 1.2, ring=mix(t.accent, t.background, 0.25))
        size = fitted(s.title, 40, 26, 40)
        title_h = text_height(s.title, size, 7.6)
        sub_h = text_height(s.subtitle, 18, 7.4) if s.subtitle else 0.0
        top = 3.6 - (title_h + sub_h) / 2
        self.text(slide, s.title, 5.0, top, 7.6, title_h, size, t.text, font=TITLE_FONT)
        if s.subtitle:
            self.text(
                slide,
                s.subtitle,
                5.0,
                top + title_h,
                7.4,
                sub_h,
                fitted(s.subtitle, 18, 13, 120),
                t.muted,
            )

    def bullets_slide(self, s: Slide) -> None:
        slide = self.new_slide(s.notes)
        top = self.heading(slide, s.title, s.subtitle)
        bottom = H - (1.6 if s.takeaway else 0.8)
        width = W - 2 * MARGIN - 0.2
        if s.icon:  # a side panel with the slide's icon
            panel = 3.3
            width -= panel + 0.4
            x = W - MARGIN - panel
            self.box(
                slide, x, top + 0.1, panel, bottom - top - 0.1, self.theme.surface, shape="round"
            )
            self.box(slide, x, top + 0.1, panel, 0.08, self.theme.accent)
            middle = top + 0.1 + (bottom - top - 0.1) / 2
            self.box(
                slide,
                x + panel / 2 - 1.15,
                middle - 1.15,
                2.3,
                2.3,
                self.theme.accent,
                shape="oval",
            )
            self.icon(
                slide, s.icon, x + panel / 2 - 1.0, middle - 1.0, 2.0, ring=self.theme.background
            )
        longest = max((len(b) for b in s.bullets), default=0)
        size = fitted("x" * longest, 22, 15, 80 if not s.icon else 60)
        self.bullet_list(slide, s.bullets, MARGIN + 0.1, top + 0.2, width, bottom - top - 0.2, size)
        self.takeaway(slide, s.takeaway)

    def cards_slide(self, s: Slide) -> None:
        """Up to four points as cards side by side: icon (or number), lead, detail."""
        slide = self.new_slide(s.notes)
        top = self.heading(slide, s.title, s.subtitle) + 0.2
        items = s.bullets[:4] or [s.subtitle]
        gap = 0.3
        width = (W - 2 * MARGIN - gap * (len(items) - 1)) / len(items)
        height = min(H - top - (1.65 if s.takeaway else 1.0), 4.4)
        details = [split_lead(i)[1] for i in items]
        size = fitted("x" * max(len(d) for d in details), 18, 12, int(width * 30))
        for i, item in enumerate(items):
            x = MARGIN + i * (width + gap)
            accent = self.theme.accent if i % 2 == 0 else self.theme.accent2
            self.box(slide, x, top, width, height, self.theme.surface, shape="round")
            self.box(slide, x, top, width, 0.08, accent)
            icon = s.icons[i] if i < len(s.icons) else ""
            if icon:
                self.icon(
                    slide,
                    icon,
                    x + 0.3,
                    top + 0.35,
                    0.95,
                    ring=mix(self.theme.surface, accent, 0.25),
                )
            else:
                self.text(
                    slide, f"{i + 1:02d}", x + 0.3, top + 0.35, 1.5, 0.8, 30, accent,
                    font=TITLE_FONT, bold=True,
                )  # fmt: skip
            lead, detail = split_lead(item)
            y = top + 1.5
            if lead:
                lead_size = fitted(lead, 20, 14, int(width * 9))
                lead_h = text_height(lead, lead_size * 1.2, width - 0.6)  # bold runs wider
                self.text(
                    slide, lead, x + 0.3, y, width - 0.6, lead_h, lead_size, self.theme.text,
                    font=TITLE_FONT, bold=True,
                )  # fmt: skip
                y += lead_h + 0.05
            self.text(
                slide, detail, x + 0.3, y, width - 0.6, top + height - y - 0.2, size,
                self.theme.text if not lead else mix(self.theme.text, self.theme.muted, 0.35),
            )  # fmt: skip
        self.takeaway(slide, s.takeaway)

    def two_column_slide(self, s: Slide) -> None:
        slide = self.new_slide(s.notes)
        top = self.heading(slide, s.title, s.subtitle) + 0.15
        gap = 0.4
        width = (W - 2 * MARGIN - gap) / 2
        height = H - top - (1.65 if s.takeaway else 0.95)
        for i, (head, items) in enumerate((s.left, s.right)):
            x = MARGIN + i * (width + gap)
            colour = self.theme.accent if i == 0 else self.theme.accent2
            self.box(slide, x, top, width, height, self.theme.surface, shape="round")
            self.text(
                slide,
                head,
                x + 0.35,
                top + 0.25,
                width - 0.7,
                0.6,
                fitted(head, 22, 16, 34),
                colour,
                font=TITLE_FONT,
            )
            longest = max((len(b) for b in items), default=0)
            self.bullet_list(
                slide,
                items,
                x + 0.35,
                top + 1.0,
                width - 0.7,
                height - 1.2,
                fitted("x" * longest, 17, 12, 48),
                marker=colour,
            )
        self.takeaway(slide, s.takeaway)

    def stats_slide(self, s: Slide) -> None:
        slide = self.new_slide(s.notes)
        top = self.heading(slide, s.title, s.subtitle) + 0.35
        stats = s.stats[:4]
        gap = 0.35
        width = (W - 2 * MARGIN - gap * (len(stats) - 1)) / max(len(stats), 1)
        height = 3.0
        icons = any(s.icons)
        if icons:
            height = 3.6
        for i, (value, label) in enumerate(stats):
            x = MARGIN + i * (width + gap)
            self.box(slide, x, top, width, height, self.theme.surface, shape="round")
            colour = self.theme.accent if i % 2 == 0 else self.theme.accent2
            self.box(slide, x, top, width, 0.08, colour)
            shift = 0.0
            if icons:
                shift = 0.75
                icon = s.icons[i] if i < len(s.icons) else ""
                self.icon(
                    slide, icon, x + width / 2 - 0.42, top + 0.3, 0.84,
                    ring=mix(self.theme.surface, colour, 0.25),
                )  # fmt: skip
            self.text(
                slide,
                value,
                x + 0.2,
                top + 0.35 + shift,
                width - 0.4,
                1.4,
                fitted(value, 54, 30, int(width * 2.2)),
                colour,
                font=TITLE_FONT,
                bold=True,
                align="center",
                anchor="middle",
            )
            self.text(
                slide,
                label,
                x + 0.3,
                top + 1.8 + shift,
                width - 0.6,
                1.1,
                fitted(label, 16, 11, int(width * 18)),
                self.theme.text,
                align="center",
            )
        if s.takeaway:
            self.takeaway(slide, s.takeaway)
        elif s.bullets:
            self.text(
                slide,
                s.bullets[0],
                MARGIN,
                top + height + 0.35,
                W - 2 * MARGIN,
                0.8,
                16,
                self.theme.muted,
                align="center",
            )

    def quote_slide(self, s: Slide) -> None:
        slide = self.new_slide(s.notes)
        t = self.theme
        self.text(slide, "“", MARGIN, 0.6, 2, 2, 140, t.accent, font=TITLE_FONT)
        quote = s.quote or s.title
        self.text(
            slide,
            quote,
            1.8,
            1.9,
            W - 3.6,
            3.2,
            fitted(quote, 34, 20, 120),
            t.text,
            font=TITLE_FONT,
            anchor="middle",
            align="center",
        )
        if s.attribution:
            self.box(slide, W / 2 - 0.5, 5.35, 1.0, 0.05, t.accent2)
            self.text(slide, s.attribution, 2, 5.5, W - 4, 0.6, 16, t.muted, align="center")

    def timeline_slide(self, s: Slide) -> None:
        slide = self.new_slide(s.notes)
        top = self.heading(slide, s.title, s.subtitle)
        steps = s.steps[:6]
        if not steps:
            return
        line_y = top + 2.0
        self.box(slide, MARGIN, line_y, W - 2 * MARGIN, 0.05, self.theme.muted)
        slot = (W - 2 * MARGIN) / len(steps)
        longest = max(len(text) for _, text in steps)
        size = fitted("x" * longest, 17, 11, int(slot * 20))
        for i, (label, text) in enumerate(steps):
            centre = MARGIN + slot * (i + 0.5)
            colour = self.theme.accent if i % 2 == 0 else self.theme.accent2
            self.box(slide, centre - 0.17, line_y - 0.15, 0.34, 0.34, colour, shape="oval")
            self.text(
                slide,
                label,
                centre - slot / 2 + 0.1,
                line_y - 1.0,
                slot - 0.2,
                0.7,
                fitted(label, 18, 12, int(slot * 9)),
                colour,
                font=TITLE_FONT,
                align="center",
                anchor="bottom",
            )
            self.text(
                slide,
                text,
                centre - slot / 2 + 0.12,
                line_y + 0.45,
                slot - 0.24,
                H - line_y - 1.4,
                size,
                self.theme.text,
                align="center",
            )

    def comparison_slide(self, s: Slide) -> None:
        from pptx.util import Inches

        slide = self.new_slide(s.notes)
        top = self.heading(slide, s.title, s.subtitle) + 0.1
        columns = s.columns[:6]
        rows = [r[: len(columns)] + [""] * (len(columns) - len(r)) for r in s.rows[:8]]
        if not columns or not rows:
            self.bullet_list(slide, s.bullets, MARGIN, top, W - 2 * MARGIN, 4, 20)
            return
        height = min(0.62 * (len(rows) + 1), H - top - (1.65 if s.takeaway else 0.9))
        shape = slide.shapes.add_table(
            len(rows) + 1,
            len(columns),
            Inches(MARGIN),
            Inches(top),
            Inches(W - 2 * MARGIN),
            Inches(height),
        )
        table = shape.table
        longest = max(len(c) for row in rows for c in row) if rows else 10
        size = fitted("x" * longest, 15, 10.5, int(70 / len(columns)))
        for c, name in enumerate(columns):
            self._cell(
                table.cell(0, c), name, self.theme.accent, self.theme.background, 15, bold=True
            )
        for r, row in enumerate(rows, 1):
            fill = self.theme.surface if r % 2 else self.theme.background
            for c, value in enumerate(row):
                colour = self.theme.text if c else self.theme.accent
                self._cell(table.cell(r, c), value, fill, colour, size, bold=c == 0)
        for row in table.rows:
            row.height = Inches(height / (len(rows) + 1))
        self.takeaway(slide, s.takeaway)

    def _cell(
        self, cell: Any, text: str, fill: str, colour: str, size: float, *, bold: bool = False
    ) -> None:
        from pptx.enum.text import MSO_ANCHOR
        from pptx.util import Inches, Pt

        cell.fill.solid()
        cell.fill.fore_color.rgb = self.rgb(fill)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.12)
        frame = cell.text_frame
        frame.word_wrap = True
        frame.text = text
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(size)
                run.font.bold = bold
                run.font.name = TITLE_FONT if bold else BODY_FONT
                run.font.color.rgb = self.rgb(colour)

    def agenda_slide(self, s: Slide) -> None:
        """The parts of the deck as numbered rows, in two columns when there are many."""
        slide = self.new_slide(s.notes)
        top = self.heading(slide, s.title or "Agenda", s.subtitle) + 0.3
        items = s.bullets[:8]
        per_column = 4 if len(items) > 4 else len(items)
        columns = math.ceil(len(items) / per_column)
        gap = 0.4
        width = (W - 2 * MARGIN - gap * (columns - 1)) / columns
        row = min(1.05, (H - top - 0.9) / per_column)
        for i, item in enumerate(items):
            x = MARGIN + (i // per_column) * (width + gap)
            y = top + (i % per_column) * row
            accent = self.theme.accent if i % 2 == 0 else self.theme.accent2
            self.box(slide, x, y, width, row - 0.15, self.theme.surface, shape="round")
            self.text(
                slide, f"{i + 1:02d}", x + 0.25, y, 1.0, row - 0.15, 26, accent,
                font=TITLE_FONT, bold=True, anchor="middle",
            )  # fmt: skip
            self.text(
                slide, item, x + 1.3, y, width - 1.5, row - 0.15,
                fitted(item, 20, 14, int(width * 7)), self.theme.text, anchor="middle",
            )  # fmt: skip

    def statement_slide(self, s: Slide) -> None:
        """One big sentence, with why it matters underneath."""
        slide = self.new_slide(s.notes)
        t = self.theme
        self.box(slide, 0, 0, 0.35, H, t.accent)
        x = MARGIN + 0.5
        width = W - 2 * x - (2.4 if s.icon else 0)
        if s.icon:
            self.icon(slide, s.icon, W - MARGIN - 2.2, H / 2 - 1.1, 2.2, ring=t.surface)
        size = fitted(s.title, 44, 30, 70)
        title_h = text_height(s.title, size, width)
        sub_h = text_height(s.subtitle, 20, width) if s.subtitle else 0.0
        top = max(0.6, 3.5 - (title_h + sub_h + 0.4) / 2)
        self.text(slide, s.title, x, top, width, title_h, size, t.text, font=TITLE_FONT)
        if s.subtitle:
            self.box(slide, x + 0.05, top + title_h + 0.1, 0.9, 0.06, t.accent2)
            self.text(
                slide, s.subtitle, x, top + title_h + 0.35, width, sub_h,
                fitted(s.subtitle, 20, 15, 160), t.muted,
            )  # fmt: skip

    def chart_slide(self, s: Slide) -> None:
        """A native PowerPoint chart (editable, data inside) with the insights beside it."""
        slide = self.new_slide(s.notes)
        top = self.heading(slide, s.title, s.subtitle) + 0.1
        bottom = H - (1.6 if s.takeaway else 0.75)
        chart_w = W - 2 * MARGIN
        if s.bullets:
            chart_w = 8.1
            x = MARGIN + chart_w + 0.35
            panel_w = W - MARGIN - x
            self.box(slide, x, top, panel_w, bottom - top, self.theme.surface, shape="round")
            self.box(slide, x, top, panel_w, 0.08, self.theme.accent2)
            self.bullet_list(
                slide, s.bullets[:3], x + 0.3, top + 0.35, panel_w - 0.55, bottom - top - 0.5,
                16, marker=self.theme.accent2, smallest=11,
            )  # fmt: skip
        if s.chart is not None:
            self.draw_chart(slide, s.chart, MARGIN, top, chart_w, bottom - top)
        self.takeaway(slide, s.takeaway)

    def draw_chart(self, slide: Any, chart: Chart, x: float, y: float, w: float, h: float) -> None:
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
        from pptx.util import Inches, Pt

        t = self.theme
        kinds = {
            "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "bar": XL_CHART_TYPE.BAR_CLUSTERED,
            "line": XL_CHART_TYPE.LINE_MARKERS,
            "pie": XL_CHART_TYPE.PIE,
            "doughnut": XL_CHART_TYPE.DOUGHNUT,
        }
        round_ = chart.kind in ("pie", "doughnut")
        data = CategoryChartData()  # type: ignore[no-untyped-call]
        data.categories = chart.categories
        for name, values in chart.series:
            data.add_series(name or "Value", values)  # type: ignore[no-untyped-call]
        frame = slide.shapes.add_chart(
            kinds[chart.kind], Inches(x), Inches(y), Inches(w), Inches(h), data
        )
        graph = frame.chart
        graph.font.size = Pt(13)
        graph.font.name = BODY_FONT
        graph.font.color.rgb = self.rgb(t.text)
        palette = [
            t.accent,
            t.accent2,
            mix(t.accent, t.text, 0.45),
            mix(t.accent2, t.background, 0.4),
        ]
        palette += [mix(t.accent, t.background, 0.45), mix(t.accent2, t.text, 0.45)]
        unit = chart.unit.replace('"', "")
        number_format = "#,##0.##"
        if unit in ("$", "£", "€", "₹", "¥"):
            number_format = f'"{unit}"{number_format}'
        elif unit == "%":
            number_format += '"%"'
        elif unit:
            number_format += f'" {unit}"'
        graph.has_title = False
        plot = graph.plots[0]
        plot.has_data_labels = round_ or len(chart.series) == 1
        if plot.has_data_labels:
            labels = plot.data_labels
            labels.number_format = number_format
            labels.number_format_is_linked = False
            labels.font.size = Pt(12)
            labels.font.color.rgb = self.rgb(t.background if round_ else t.text)
        if round_:
            graph.has_legend = True
            graph.legend.position = XL_LEGEND_POSITION.RIGHT
            graph.legend.include_in_layout = False
            graph.legend.font.size = Pt(13)
            for i, point in enumerate(plot.series[0].points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = self.rgb(palette[i % len(palette)])
                point.format.line.color.rgb = self.rgb(t.background)
            return
        graph.has_legend = len(chart.series) > 1
        if graph.has_legend:
            graph.legend.position = XL_LEGEND_POSITION.TOP
            graph.legend.include_in_layout = False
        if chart.kind in ("column", "bar"):
            plot.gap_width = 70
            plot.overlap = -10 if len(chart.series) > 1 else 0
        for i, series in enumerate(plot.series):
            colour = self.rgb(palette[i % len(palette)])
            if chart.kind == "line":
                series.format.line.color.rgb = colour
                series.format.line.width = Pt(3)
                series.smooth = False
                series.marker.format.fill.solid()
                series.marker.format.fill.fore_color.rgb = colour
            else:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = colour
        values = graph.value_axis
        values.has_major_gridlines = True
        values.major_gridlines.format.line.color.rgb = self.rgb(mix(t.background, t.muted, 0.35))
        values.format.line.fill.background()
        values.tick_labels.number_format = number_format
        values.tick_labels.number_format_is_linked = False
        values.tick_labels.font.color.rgb = self.rgb(t.muted)
        categories = graph.category_axis
        categories.format.line.color.rgb = self.rgb(t.muted)
        categories.tick_labels.font.color.rgb = self.rgb(t.text)
        if chart.kind == "bar":  # first category on top, as it reads in the list
            categories.reverse_order = True

    def closing_slide(self, s: Slide) -> None:
        slide = self.new_slide(s.notes, footer=False)
        t = self.theme
        self.box(slide, -1.5, H - 3.2, 5, 5, t.surface, shape="oval")
        self.box(slide, W - 3.3, -1.2, 4.2, 4.2, t.accent, shape="oval")
        title = s.title or "Thank you"
        self.text(
            slide,
            title,
            MARGIN,
            1.2,
            9,
            1.4,
            fitted(title, 44, 28, 34),
            t.text,
            font=TITLE_FONT,
            anchor="bottom",
        )
        self.box(slide, MARGIN + 0.05, 2.75, 0.9, 0.06, t.accent2)
        if s.bullets:
            self.bullet_list(slide, s.bullets[:5], MARGIN, 3.15, 8.6, 3.4, 20)
        elif s.subtitle:
            self.text(slide, s.subtitle, MARGIN, 3.1, 8.6, 1.2, 20, t.muted)


def render(deck: Deck, path: Path) -> int:
    """Write ``deck`` to ``path`` (.pptx). Returns the number of slides."""
    theme = THEMES.get(deck.theme, THEMES["midnight"])
    painter = _Painter(theme, deck.title)
    section = 0
    for slide in deck.slides:
        layout = slide.layout if slide.layout in LAYOUTS else "bullets"
        if layout == "section":
            section += 1
            painter.section_slide(slide, section)
        elif layout == "title":
            painter.title_slide(slide)
        elif layout == "closing":
            painter.closing_slide(slide)
        else:
            getattr(painter, f"{layout}_slide")(slide)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("xb") as handle:  # never overwrites
        painter.prs.save(handle)
    return painter.number
